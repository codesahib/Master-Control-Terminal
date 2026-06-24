import re
from datetime import date, datetime
from decimal import Decimal

import openpyxl
from sqlalchemy import delete, select

from app.db.session import SessionLocal
from app.models.models import (
    Account,
    Category,
    ContributionLimit,
    HoldingSnapshot,
    Instrument,
    Platform,
    Transaction,
    TransactionType,
)
from app.services.finance import get_contribution_room, normalize_platform, seed_reference_data

FILE_PATH = r"C:/Users/gursa/Downloads/Untitled spreadsheet.xlsx"
YEAR_SHEETS = ["2023", "2024", "2025", "2026"]
CURRENT_YEAR = 2026
SNAPSHOT_DATE = date.today()
IMPORT_CONTRIBUTIONS = False


def to_float(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, Decimal):
        return float(v)
    if isinstance(v, str):
        s = v.strip().replace(",", "")
        if not s:
            return None
        try:
            return float(s)
        except ValueError:
            return None
    return None


def guess_platform(label: str):
    text = label.lower()
    if "ws" in text or "wealthsimple" in text:
        return "Wealthsimple"
    if "qt" in text or "questrade" in text:
        return "Questrade"
    if "sunlife" in text or "sun life" in text:
        return "Sun Life"
    if "cibc" in text:
        return "CIBC"
    return None


def get_or_create_account(db, name: str):
    acc = db.scalar(select(Account).where(Account.name == name))
    if not acc:
        acc = Account(name=name)
        db.add(acc)
        db.flush()
    return acc


def get_or_create_instrument(db, symbol: str):
    sym = symbol.strip().upper()[:50]
    inst = db.scalar(select(Instrument).where(Instrument.symbol == sym))
    if not inst:
        inst = Instrument(symbol=sym, name=symbol.strip())
        db.add(inst)
        db.flush()
    return inst


def get_or_create_category(db, broad: str, precise: str):
    broad = (broad or "Unknown").strip()[:100]
    precise = (precise or "Unknown").strip()[:100]
    cat = db.scalar(select(Category).where(Category.broad == broad, Category.precise == precise))
    if not cat:
        cat = Category(broad=broad, precise=precise)
        db.add(cat)
        db.flush()
    return cat


def holding_record_type(fund: str | None, broad: str | None) -> str:
    if (fund or "").strip().lower() == "unused":
        return "unused"
    if (broad or "").strip().lower() == "cash":
        return "cash"
    return "holding"


def current_year_unused_by_account(db, tax_year: int) -> dict[str, float]:
    rows = get_contribution_room(db, tax_year)
    return {row["account"]: float(row["remaining"]) for row in rows}


def main():
    wb = openpyxl.load_workbook(FILE_PATH, data_only=True)
    db = SessionLocal()
    seed_reference_data(db)

    # Replace only the workbook-driven current holdings snapshot.
    db.execute(
        delete(HoldingSnapshot).where(
            HoldingSnapshot.snapshot_year == CURRENT_YEAR,
            HoldingSnapshot.snapshot_type == "current",
        )
    )
    db.commit()

    inserted_limits = 0
    inserted_txns = 0
    inserted_holdings = 0

    # Contribution imports are intentionally opt-in; this task imports holdings only.
    for sheet_name in YEAR_SHEETS if IMPORT_CONTRIBUTIONS else []:
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        tax_year = int(sheet_name)

        for row, account_name in [(7, "RRSP"), (8, "FHSA"), (9, "TFSA")]:
            account = get_or_create_account(db, account_name)
            new = to_float(ws.cell(row, 11).value) or 0.0
            if new == 0 and tax_year == 2023:
                new = to_float(ws.cell(row, 12).value) or 0.0
            existing_limit = db.scalar(
                select(ContributionLimit).where(
                    ContributionLimit.account_id == account.id,
                    ContributionLimit.tax_year == str(tax_year),
                )
            )
            if not existing_limit:
                db.add(
                    ContributionLimit(
                        account_id=account.id,
                        tax_year=str(tax_year),
                        new_room=new,
                    )
                )
                inserted_limits += 1

        # RRSP/TFSA block I:L rows 12-23
        for r in range(12, 24):
            for account_name, label_col, amount_col in [("RRSP", 9, 10), ("TFSA", 11, 12)]:
                label = ws.cell(r, label_col).value
                amount = to_float(ws.cell(r, amount_col).value)
                if not label or amount is None or amount <= 0:
                    continue
                account = get_or_create_account(db, account_name)
                platform_name = guess_platform(str(label))
                platform = normalize_platform(db, platform_name) if platform_name else None
                exists = db.scalar(
                    select(Transaction).where(
                        Transaction.transaction_type == TransactionType.contribution,
                        Transaction.transaction_date == date(tax_year, 12, 31),
                        Transaction.account_id == account.id,
                        Transaction.amount == amount,
                        Transaction.notes == str(label),
                    )
                )
                if exists:
                    continue
                db.add(
                    Transaction(
                        transaction_type=TransactionType.contribution,
                        transaction_date=date(tax_year, 12, 31),
                        account_id=account.id,
                        platform_id=platform.id if platform else None,
                        amount=amount,
                        fees=0,
                        notes=str(label),
                    )
                )
                inserted_txns += 1

        # FHSA block M:N rows 12-23
        for r in range(12, 24):
            label = ws.cell(r, 13).value
            amount = to_float(ws.cell(r, 14).value)
            if not label or amount is None or amount <= 0:
                continue
            account = get_or_create_account(db, "FHSA")
            platform_name = guess_platform(str(label))
            platform = normalize_platform(db, platform_name) if platform_name else None
            exists = db.scalar(
                select(Transaction).where(
                    Transaction.transaction_type == TransactionType.contribution,
                    Transaction.transaction_date == date(tax_year, 12, 31),
                    Transaction.account_id == account.id,
                    Transaction.amount == amount,
                    Transaction.notes == str(label),
                )
            )
            if exists:
                continue
            db.add(
                Transaction(
                    transaction_type=TransactionType.contribution,
                    transaction_date=date(tax_year, 12, 31),
                    account_id=account.id,
                    platform_id=platform.id if platform else None,
                    amount=amount,
                    fees=0,
                    notes=str(label),
                )
            )
            inserted_txns += 1

    db.commit()

    # Holdings from Distrubution/Distribution sheet.
    dist_name = "Distrubution" if "Distrubution" in wb.sheetnames else "Distribution"
    if dist_name in wb.sheetnames:
        ws = wb[dist_name]
        unused_by_account = current_year_unused_by_account(db, CURRENT_YEAR)
        for r in range(17, 42):
            for start_col, account_name in [(1, "FHSA"), (9, "RRSP"), (16, "TFSA")]:
                fund = ws.cell(r, start_col).value
                broad = ws.cell(r, start_col + 1).value
                precise = ws.cell(r, start_col + 2).value
                amount = to_float(ws.cell(r, start_col + 3).value)
                if not fund or amount is None:
                    continue
                record_type = holding_record_type(str(fund), str(broad) if broad else None)
                if record_type == "unused":
                    amount = unused_by_account.get(account_name, amount)
                account = get_or_create_account(db, account_name)
                instrument = get_or_create_instrument(db, str(fund))
                category = get_or_create_category(db, str(broad) if broad else "Unknown", str(precise) if precise else "Unknown")
                db.add(
                    HoldingSnapshot(
                        snapshot_date=SNAPSHOT_DATE,
                        snapshot_year=CURRENT_YEAR,
                        snapshot_type="current",
                        holding_date=None,
                        record_type=record_type,
                        account_id=account.id,
                        instrument_id=instrument.id,
                        category_id=category.id,
                        market_value=amount,
                    )
                )
                inserted_holdings += 1

    db.commit()
    print({
        "limits_inserted": inserted_limits,
        "transactions_inserted": inserted_txns,
        "holdings_inserted": inserted_holdings,
    })


if __name__ == "__main__":
    main()
