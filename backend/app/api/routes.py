import json
from datetime import datetime, timezone
from typing import Literal

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from fastapi.encoders import jsonable_encoder
from fastapi.responses import JSONResponse
from sqlalchemy import select
from sqlalchemy.exc import SQLAlchemyError
from sqlalchemy.orm import Session
from openpyxl import load_workbook

from app.db.session import get_db
from app.models.models import ImportType, Platform, Transaction, TransactionType
from app.schemas.schemas import (
    AccountTransactionCreate,
    AccountTransactionRead,
    PaginatedAccountTransactionRead,
    ContributionCreate,
    ContributionFundingRead,
    ContributionLimitRead,
    ContributionLimitUpdate,
    ContributionRead,
    ContributionRoomRead,
    DistributionPoint,
    HoldingRead,
    ImportPreviewResponse,
    ImportPreviewRow,
    MarketPriceRefreshResult,
    ManualValuationUpsert,
    PaginatedTransactionRead,
    PaginatedPortfolioPLRead,
    PortfolioSummaryRead,
    SymbolRead,
    TimeSeriesPoint,
    TransactionCreate,
    TransactionRead,
)
from app.services.finance import (
    create_account_transaction,
    create_contribution,
    create_import_preview,
    create_transaction,
    export_all_data,
    get_all_contribution_room,
    get_contribution_room,
    list_available_contributions,
    list_contribution_limits,
    restore_all_data,
    TRACKED_YEARS,
    upsert_manual_valuation,
    upsert_contribution_limit,
    update_account_transaction,
    update_contribution,
    update_transaction,
)
from app.services.market_data import refresh_market_prices, search_symbols
from app.services.portfolio import distribution, grouped_holdings, portfolio_pl, portfolio_summary, timeseries
from app.services.transaction_reads import (
    account_transaction_reads,
    contribution_read,
    contribution_reads,
    transaction_read,
    transaction_reads,
)

router = APIRouter()


@router.get("/years", response_model=list[int])
def list_years():
    return [int(year) for year in TRACKED_YEARS]


@router.get("/platforms", response_model=list[str])
def list_platforms(db: Session = Depends(get_db)):
    return list(db.scalars(select(Platform.canonical_name).order_by(Platform.canonical_name)))


@router.get("/symbols/search", response_model=list[SymbolRead])
def search_symbols_endpoint(
    q: str = Query(min_length=1),
    limit: int = Query(default=8, ge=1, le=20),
    db: Session = Depends(get_db),
):
    return search_symbols(db, q, limit)


@router.post("/market-prices/refresh", response_model=list[MarketPriceRefreshResult])
def refresh_market_prices_endpoint(db: Session = Depends(get_db)):
    return refresh_market_prices(db)


@router.post("/manual-valuations")
def upsert_manual_valuation_endpoint(payload: ManualValuationUpsert, db: Session = Depends(get_db)):
    try:
        snapshot = upsert_manual_valuation(db, payload)
    except ValueError as exc:
        db.rollback()
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    return {
        "id": snapshot.id,
        "snapshot_date": snapshot.snapshot_date,
        "market_value": float(snapshot.market_value),
    }


@router.get("/portfolio/pl", response_model=PaginatedPortfolioPLRead)
def portfolio_pl_endpoint(
    account: str | None = Query(default=None),
    platform: str | None = Query(default=None),
    status: str | None = Query(default=None),
    year: int | None = Query(default=None),
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=10, ge=1, le=500),
    sort_direction: Literal["asc", "desc"] = Query(default="desc"),
    db: Session = Depends(get_db),
):
    rows = portfolio_pl(db, account=account, platform=platform, status=status, year=year)
    rows = sorted(rows, key=lambda row: row["symbol"].lower(), reverse=sort_direction == "desc")
    total = len(rows)
    start = (page - 1) * page_size
    return PaginatedPortfolioPLRead(items=rows[start:start + page_size], total=total, page=page, page_size=page_size)


@router.get("/portfolio/summary", response_model=PortfolioSummaryRead)
def portfolio_summary_endpoint(
    year: int | None = Query(default=None),
    db: Session = Depends(get_db),
):
    return portfolio_summary(db, year=year)


@router.post("/transactions", response_model=TransactionRead)
def create_transaction_endpoint(payload: TransactionCreate, db: Session = Depends(get_db)):
    try:
        txn = create_transaction(db, payload)
    except ValueError as exc:
        db.rollback()
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    return transaction_read(db, txn)


@router.put("/transactions/{transaction_id}", response_model=TransactionRead)
def update_transaction_endpoint(transaction_id: int, payload: TransactionCreate, db: Session = Depends(get_db)):
    try:
        txn = update_transaction(db, transaction_id, payload)
    except ValueError as exc:
        db.rollback()
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    if not txn:
        raise HTTPException(status_code=404, detail="Transaction not found")
    return transaction_read(db, txn)


@router.get("/transactions", response_model=PaginatedTransactionRead)
def list_transactions_endpoint(
    transaction_type: TransactionType | None = Query(default=None),
    account: str | None = Query(default=None),
    platform: str | None = Query(default=None),
    symbol: str | None = Query(default=None),
    year: int | None = Query(default=None),
    page: int = Query(default=1, ge=1),
    sort_direction: Literal["asc", "desc"] = Query(default="desc"),
    db: Session = Depends(get_db),
):
    items, total = transaction_reads(
        db, transaction_type=transaction_type, account=account, platform=platform, symbol=symbol,
        year=year, page=page, sort_direction=sort_direction,
    )
    return PaginatedTransactionRead(
        items=items, total=total, page=page, page_size=10
    )


@router.post("/contributions", response_model=ContributionRead)
def create_contribution_endpoint(payload: ContributionCreate, db: Session = Depends(get_db)):
    txn = create_contribution(db, payload)
    return contribution_read(db, txn)


@router.put("/contributions/{transaction_id}", response_model=ContributionRead)
def update_contribution_endpoint(transaction_id: int, payload: ContributionCreate, db: Session = Depends(get_db)):
    txn = update_contribution(db, transaction_id, payload)
    if not txn:
        raise HTTPException(status_code=404, detail="Contribution not found")
    return contribution_read(db, txn)


@router.get("/contributions", response_model=list[ContributionRead])
def list_contributions_endpoint(
    account: str | None = Query(default=None),
    platform: str | None = Query(default=None),
    year: int | None = Query(default=None),
    db: Session = Depends(get_db),
):
    return contribution_reads(db, account=account, platform=platform, year=year)


@router.get("/available-contributions", response_model=list[ContributionFundingRead])
def list_available_contributions_endpoint(
    account: str,
    exclude_transaction_id: int | None = None,
    db: Session = Depends(get_db),
):
    return list_available_contributions(db, account, db.get(Transaction, exclude_transaction_id) if exclude_transaction_id else None)


@router.post("/account-transactions", response_model=AccountTransactionRead)
def create_account_transaction_endpoint(payload: AccountTransactionCreate, db: Session = Depends(get_db)):
    try:
        txn = create_account_transaction(db, payload)
    except ValueError as exc:
        db.rollback()
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    return transaction_read(db, txn)


@router.put("/account-transactions/{transaction_id}", response_model=AccountTransactionRead)
def update_account_transaction_endpoint(transaction_id: int, payload: AccountTransactionCreate, db: Session = Depends(get_db)):
    try:
        txn = update_account_transaction(db, transaction_id, payload)
    except ValueError as exc:
        db.rollback()
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    if not txn:
        raise HTTPException(status_code=404, detail="Account transaction not found")
    return transaction_read(db, txn)


@router.get("/account-transactions", response_model=PaginatedAccountTransactionRead)
def list_account_transactions_endpoint(
    account: str | None = Query(default=None),
    platform: str | None = Query(default=None),
    symbol: str | None = Query(default=None),
    year: int | None = Query(default=None),
    transaction_type: TransactionType | None = Query(default=None),
    page: int = Query(default=1, ge=1),
    sort_direction: Literal["asc", "desc"] = Query(default="desc"),
    db: Session = Depends(get_db),
):
    items, total = account_transaction_reads(
        db, account=account, platform=platform, symbol=symbol, year=year,
        transaction_type=transaction_type, page=page, sort_direction=sort_direction,
    )
    return PaginatedAccountTransactionRead(
        items=items, total=total, page=page, page_size=10
    )


@router.get("/holdings", response_model=list[HoldingRead])
def list_holdings_endpoint(
    account: str | None = Query(default=None),
    year: int | None = Query(default=None),
    db: Session = Depends(get_db),
):
    return grouped_holdings(db, account=account, year=year)


@router.get("/limits/{year}", response_model=list[ContributionRoomRead])
def get_limits(year: int, db: Session = Depends(get_db)):
    return get_contribution_room(db, year)


@router.get("/limits", response_model=list[ContributionRoomRead])
def get_all_limits(db: Session = Depends(get_db)):
    return get_all_contribution_room(db)


@router.get("/contribution-limits", response_model=list[ContributionLimitRead])
def get_contribution_limits(db: Session = Depends(get_db)):
    return list_contribution_limits(db)


@router.put("/contribution-limits/{account}/{tax_year}", response_model=ContributionLimitRead)
def update_contribution_limit(
    account: str,
    tax_year: str,
    payload: ContributionLimitUpdate,
    db: Session = Depends(get_db),
):
    if not tax_year.isdigit() or len(tax_year) != 4:
        raise HTTPException(status_code=400, detail="tax_year must be a 4-digit year")
    return upsert_contribution_limit(
        db,
        account,
        tax_year,
        payload.new_room,
    )


@router.get("/analytics/distribution", response_model=list[DistributionPoint])
def get_distribution(
    group_by: str = Query(pattern="^(sector|account|platform)$"),
    year: int | None = Query(default=None),
    category_level: str = Query(default="precise", pattern="^(broad|precise)$"),
    currency: str = Query(default="CAD", pattern="^[A-Z]{3}$"),
    db: Session = Depends(get_db),
):
    return distribution(db, group_by, year, category_level, currency)


@router.get("/analytics/timeseries", response_model=list[TimeSeriesPoint])
def get_timeseries(
    year: int | None = Query(default=None),
    currency: str = Query(default="CAD", pattern="^[A-Z]{3}$"),
    db: Session = Depends(get_db),
):
    return timeseries(db, year, currency)


@router.get("/export")
def export_data(db: Session = Depends(get_db)):
    exported_at = datetime.now(timezone.utc)
    payload = {
        "meta": {
            "exported_at": exported_at,
            "format": "master-control-terminal-export",
            "version": 2,
        },
        "data": export_all_data(db),
    }
    filename = f"master-control-terminal-export-{exported_at.strftime('%Y%m%dT%H%M%SZ')}.json"
    return JSONResponse(
        content=jsonable_encoder(payload),
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.post("/import-backup")
async def import_backup(file: UploadFile = File(...), db: Session = Depends(get_db)):
    if not file.filename or not file.filename.lower().endswith(".json"):
        raise HTTPException(status_code=400, detail="Please upload a JSON backup file")

    try:
        payload = json.loads(await file.read())
    except json.JSONDecodeError as exc:
        raise HTTPException(status_code=400, detail="Backup file is not valid JSON") from exc

    meta = payload.get("meta", {})
    if meta.get("format") != "master-control-terminal-export":
        raise HTTPException(status_code=400, detail="Unsupported backup format")

    try:
        restore_all_data(db, payload)
    except (ValueError, SQLAlchemyError) as exc:
        db.rollback()
        raise HTTPException(status_code=400, detail=str(exc)) from exc

    return {"status": "ok", "message": "Backup restored successfully"}


@router.post("/imports/contributions", response_model=ImportPreviewResponse)
async def import_contributions(file: UploadFile = File(...), db: Session = Depends(get_db)):
    wb = load_workbook(file.file, data_only=True)
    rows = []
    for year in [name for name in wb.sheetnames if name.isdigit()]:
        ws = wb[year]
        for r in range(12, 24):
            rrsp_label = ws.cell(r, 9).value
            rrsp_amount = ws.cell(r, 10).value
            tfsa_label = ws.cell(r, 11).value
            tfsa_amount = ws.cell(r, 12).value
            if rrsp_label and rrsp_amount:
                rows.append({"tax_year": int(year), "account": "RRSP", "label": rrsp_label, "amount": float(rrsp_amount)})
            if tfsa_label and tfsa_amount:
                rows.append({"tax_year": int(year), "account": "TFSA", "label": tfsa_label, "amount": float(tfsa_amount)})

    imp = create_import_preview(db, ImportType.contributions, file.filename, rows)
    return ImportPreviewResponse(
        import_id=imp.id,
        rows=[ImportPreviewRow(row_number=i + 1, payload=row) for i, row in enumerate(rows[:250])],
        row_count=len(rows),
    )


@router.post("/imports/holdings", response_model=ImportPreviewResponse)
async def import_holdings(file: UploadFile = File(...), db: Session = Depends(get_db)):
    wb = load_workbook(file.file, data_only=True)
    if "Distrubution" not in wb.sheetnames and "Distribution" not in wb.sheetnames:
        raise HTTPException(status_code=400, detail="Distribution sheet not found")

    sheet_name = "Distrubution" if "Distrubution" in wb.sheetnames else "Distribution"
    ws = wb[sheet_name]
    rows = []

    for r in range(17, 42):
        for start_col, account in [(1, "FHSA"), (9, "RRSP"), (16, "TFSA")]:
            fund = ws.cell(r, start_col).value
            broad = ws.cell(r, start_col + 1).value
            precise = ws.cell(r, start_col + 2).value
            amount = ws.cell(r, start_col + 3).value
            if fund and amount:
                rows.append(
                    {
                        "account": account,
                        "instrument": str(fund),
                        "broad_category": str(broad) if broad else "Unknown",
                        "precise_category": str(precise) if precise else "Unknown",
                        "market_value": float(amount),
                    }
                )

    imp = create_import_preview(db, ImportType.holdings, file.filename, rows)
    return ImportPreviewResponse(
        import_id=imp.id,
        rows=[ImportPreviewRow(row_number=i + 1, payload=row) for i, row in enumerate(rows[:250])],
        row_count=len(rows),
    )
